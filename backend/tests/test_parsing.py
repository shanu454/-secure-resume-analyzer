import pytest

from app.core.config import get_settings
from app.services.parsing import DocumentParseError, extract_text, normalize_text
from tests.conftest import RESUME_LINES, make_docx_bytes, make_pdf_bytes

S = get_settings()


def test_extract_from_pdf():
    text = extract_text(make_pdf_bytes(), "pdf", S)
    assert "FastAPI" in text and "PostgreSQL" in text


def test_extract_from_docx_including_tables():
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("Experience with Python and Docker for backend services.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Kubernetes"
    table.rows[0].cells[1].text = "Terraform"
    buf = BytesIO()
    doc.save(buf)
    text = extract_text(buf.getvalue(), "docx", S)
    assert "Python" in text and "Kubernetes | Terraform" in text


def test_extract_from_txt_with_bom():
    text = extract_text("\ufeffPython developer with Docker experience".encode(), "txt", S)
    assert text.startswith("Python")


def test_corrupt_files_give_generic_error():
    with pytest.raises(DocumentParseError) as exc:
        extract_text(b"%PDF-1.4 this is not really a pdf SECRET-CONTENT", "pdf", S)
    assert "SECRET-CONTENT" not in str(exc.value)


def test_image_only_documents_rejected():
    from io import BytesIO

    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf)
    c.rect(10, 10, 100, 100)
    c.save()
    with pytest.raises(DocumentParseError, match="No readable text"):
        extract_text(buf.getvalue(), "pdf", S)


def test_too_many_pages_rejected():
    from io import BytesIO

    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf)
    for i in range(S.max_pdf_pages + 1):
        c.drawString(50, 800, f"Page {i} with enough readable text to count")
        c.showPage()
    c.save()
    with pytest.raises(DocumentParseError, match="pages"):
        extract_text(buf.getvalue(), "pdf", S)


def test_too_much_text_rejected():
    with pytest.raises(DocumentParseError, match="too much text"):
        extract_text(("python " * 100_000).encode(), "txt", S)


def test_normalize_text():
    assert normalize_text("a\x00b\r\n\r\n\r\n\r\nc   d") == "a b\n\nc d"


def test_docx_helper_roundtrip():
    assert "Backend developer" in extract_text(make_docx_bytes(RESUME_LINES), "docx", S)
