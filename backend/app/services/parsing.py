"""Extract plain text from validated PDF, DOCX and TXT files."""
import io
import re
import unicodedata

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.config import Settings

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class DocumentParseError(Exception):
    """Raised with a safe, generic message. Never includes document content."""


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = _CONTROL.sub(" ", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes, settings: Settings) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise DocumentParseError("Password-protected PDFs are not supported.")
    if len(reader.pages) > settings.max_pdf_pages:
        raise DocumentParseError(f"PDF has more than {settings.max_pdf_pages} pages.")
    parts = []
    size = 0
    for page in reader.pages:
        chunk = page.extract_text() or ""
        size += len(chunk)
        if size > settings.max_text_chars:
            raise DocumentParseError("Document contains too much text.")
        parts.append(chunk)
    return "\n".join(parts)


def _docx_text(data: bytes, settings: Settings) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, file_type: str, settings: Settings) -> str:
    try:
        if file_type == "pdf":
            raw = _pdf_text(data, settings)
        elif file_type == "docx":
            raw = _docx_text(data, settings)
        elif file_type == "txt":
            raw = data.decode("utf-8-sig")
        else:
            raise DocumentParseError("Unsupported file type.")
    except DocumentParseError:
        raise
    except Exception:  # parser libraries raise many types; never leak their messages
        raise DocumentParseError("The file could not be read. It may be corrupted.") from None

    text = normalize_text(raw)
    if len(text) > settings.max_text_chars:
        raise DocumentParseError("Document contains too much text.")
    if len(text) < 20:
        raise DocumentParseError(
            "No readable text was found. Scanned or image-only documents are not supported."
        )
    return text
